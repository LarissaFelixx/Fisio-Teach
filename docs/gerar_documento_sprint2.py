# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()

# --- Página A4 e margens padrão do Word ---
sec = DOC.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(3.0)

# --- Fonte base: Calibri 11 (mesma do documento da Sprint 1) ---
normal = DOC.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15


def p(text='', bold=False, italic=False, size=11, align=None, space_after=8,
      space_before=0, color=None, font='Calibri'):
    par = DOC.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        par.alignment = align
    if text:
        run = par.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = font
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
        if color:
            run.font.color.rgb = color
    return par


def heading(text):
    return p(text, bold=True, size=14, space_before=14, space_after=6,
             color=RGBColor(0x1F, 0x38, 0x64))


def hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    for tag, val in (('w:color', '0563C1'), ('w:u', 'single')):
        el = OxmlElement(tag)
        el.set(qn('w:val'), 'single' if tag == 'w:u' else val)
        rPr.append(el)
    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:ascii'), 'Calibri')
    rfonts.set(qn('w:hAnsi'), 'Calibri')
    rPr.append(rfonts)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=11, align=None):
    cell.text = ''
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.space_before = Pt(2)
    if align is not None:
        par.alignment = align
    run = par.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')



def fix_widths(table, widths):
    """Força layout fixo: sem isso o Word/LibreOffice distribui as colunas igualmente."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(int(Cm(w).twips)))
        grid.append(col)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def build_table(headers, rows, widths):
    table = DOC.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell_text(cell, h, bold=True)
        shade(cell, 'DEEAF6')
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cell_text(cells[i], v)
    fix_widths(table, widths)
    return table


# ============================ CAPA ============================
for line in ('UNIVERSIDADE DE FORTALEZA',
             'CENTRO DE CIÊNCIAS TECNOLÓGICAS',
             'TECNÓLOGO EM ANÁLISE E DESENVOLVIMENTO DE SISTEMAS',
             'N393 — Projeto Aplicado Multiplataforma'):
    p(line, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

p('', space_after=0)
p('DOCUMENTAÇÃO DE ENTREGA', bold=True, size=14,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=0)
p('Sprint 2', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
p('', space_after=0)
p('Fortaleza — 2025', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18)

# ===================== 1. NOME DO PROJETO =====================
heading('1. Nome do projeto')
p('FisioTech — aplicativo mobile de gestão para clínicas e consultórios de fisioterapia, '
  'reunindo cadastro de pacientes e profissionais, agendamento e acompanhamento clínico de '
  'consultas e comunicação entre paciente e profissional em uma única base de código, '
  'distribuída como aplicação web e como aplicativo Android nativo.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)
p('Supervisora: Profª. Lyndainês Santos.')

# ======================= 2. INTEGRANTES =======================
heading('2. Integrantes')
build_table(
    ['Nome do integrante', 'Matrícula'],
    [['Gabriel Raniere Barbosa Neri', '2412891'],
     ['Pedro Italo do Nascimento Costa', '2426040'],
     ['Marco Antonio Pessoa Noronha', '2427661'],
     ['Larissa Maria Félix de Oliveira', '2422783'],
     ['Caio Gabriel Rodrigues de Melo', '2422764']],
    [10.0, 5.0])

# ==================== 3. LINKS DO PROJETO =====================
heading('3. Links do projeto')
p('Repositório de código e quadro de acompanhamento da equipe.')
links = DOC.add_table(rows=1, cols=2)
links.style = 'Table Grid'
links.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(['Recurso', 'Endereço']):
    cell_text(links.rows[0].cells[i], h, bold=True)
    shade(links.rows[0].cells[i], 'DEEAF6')
for nome, url in [('Repositório', 'https://github.com/gabrielneriqa/fisiotech-back'),
                  ('Project (board)', 'https://github.com/gabrielneriqa/fisiotech-back/projects')]:
    cells = links.add_row().cells
    cell_text(cells[0], nome)
    cells[1].text = ''
    par = cells[1].paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.space_before = Pt(2)
    hyperlink(par, url, 'Link')
fix_widths(links, [6.0, 9.0])

# ================= 4. ATIVIDADES DA SPRINT ====================
heading('4. Atividades da Sprint')
p('Com a fundação técnica estabelecida na Sprint 1, a Sprint 2 desloca o foco da construção '
  'inicial para o amadurecimento do produto. O trabalho se organiza em três frentes '
  'complementares: elevar o nível de segurança e de confiabilidade do que já existe '
  '(autenticação por token, versionamento do banco, testes de integração e integração '
  'contínua), tirar a aplicação do ambiente local e colocá-la em um ambiente publicado com o '
  'aplicativo Android conectado a ele, e ampliar o valor clínico do sistema com o '
  'acompanhamento da evolução do paciente e a agenda do profissional.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)

build_table(
    ['Item', 'Atividade', 'Status'],
    [['Fisio-Teach #5', 'Migrar a autenticação de HTTP Basic para JWT', 'Backlog'],
     ['Fisio-Teach #6', 'Versionar o schema do banco com migrações Flyway', 'Backlog'],
     ['Fisio-Teach #7', 'Publicar a API em ambiente de homologação e gerar o .apk conectado a ela', 'Backlog'],
     ['Fisio-Teach #8', 'Automatizar build e testes em pipeline de integração contínua', 'Backlog'],
     ['Fisio-Teach #9', 'Implementar a recuperação de senha por e-mail', 'Backlog'],
     ['Fisio-Teach #10', 'Registrar a evolução clínica e o plano terapêutico por sessão', 'Backlog'],
     ['Fisio-Teach #11', 'Construir a agenda e os indicadores do profissional', 'Backlog'],
     ['Fisio-Teach #12', 'Paginar, filtrar e ordenar as listagens da API e do aplicativo', 'Backlog'],
     ['Fisio-Teach #13', 'Cobrir os endpoints com testes de integração automatizados', 'Backlog'],
     ['Fisio-Teach #14', 'Criação de documentação e entrega de sprint', 'Backlog']],
    [3.4, 9.2, 2.4])

p('Registro do quadro de atividades:', space_before=12)
ph = p('[Inserir aqui a captura de tela do quadro da Sprint 2 no project do repositório]',
       italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       color=RGBColor(0x80, 0x80, 0x80), space_before=6, space_after=6)
p('Figura 1 — Backlog da Sprint 2 no project do repositório.',
  italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

out = '/home/user/fisiotech-back/docs/FisioTech_Documentacao_Entrega_Sprint2.docx'
DOC.save(out)
print('salvo:', out)
