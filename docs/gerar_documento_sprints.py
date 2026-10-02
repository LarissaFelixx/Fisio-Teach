# -*- coding: utf-8 -*-
"""Gera a documentação consolidada de entrega das Sprints 1 a 4 do FisioTech.

As atividades foram reconstruídas a partir do histórico de commits do repositório
(`git log --reverse master`), de modo que o quadro de sprints reflita o que já foi
efetivamente entregue.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()

sec = DOC.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(3.0)

normal = DOC.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15


def p(text='', bold=False, italic=False, size=11, align=None, space_after=8,
      space_before=0, color=None):
    par = DOC.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        par.alignment = align
    if text:
        run = par.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(size)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        if color:
            run.font.color.rgb = color
    return par


def heading(text, page_break=False):
    if page_break:
        brk = DOC.add_paragraph()
        brk.paragraph_format.space_after = Pt(0)
        brk.add_run().add_break(WD_BREAK.PAGE)
    par = p(text, bold=True, size=14, space_before=0 if page_break else 14,
            space_after=6, color=RGBColor(0x1F, 0x38, 0x64))
    par.paragraph_format.keep_with_next = True
    return par


def hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), 'Calibri'); rf.set(qn('w:hAnsi'), 'Calibri'); rPr.append(rf)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=11):
    cell.text = ''
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.space_before = Pt(2)
    run = par.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')


def fix_widths(table, widths):
    """Layout fixo: sem isso Word e LibreOffice distribuem as colunas igualmente."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(int(Cm(w).twips)))
        grid.append(col)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def build_table(headers, rows, widths, size=11):
    table = DOC.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True, size=size)
        shade(table.rows[0].cells[i], 'DEEAF6')
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cell_text(cells[i], v, size=size)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for row in table.rows:
        cant = OxmlElement('w:cantSplit'); cant.set(qn('w:val'), 'true')
        row._tr.get_or_add_trPr().append(cant)
    fix_widths(table, widths)
    return table


def figura(numero, sprint):
    p('Registro do quadro de atividades:', space_before=12)
    p(f'[Inserir aqui a captura de tela da Sprint {sprint} no project do repositório]',
      italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
      color=RGBColor(0x80, 0x80, 0x80), space_before=6, space_after=6)
    p(f'Figura {numero} — Quadro da Sprint {sprint} no project do repositório.',
      italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)


# ============================== CAPA ==============================
for line in ('UNIVERSIDADE DE FORTALEZA',
             'CENTRO DE CIÊNCIAS TECNOLÓGICAS',
             'TECNÓLOGO EM ANÁLISE E DESENVOLVIMENTO DE SISTEMAS',
             'N393 — Projeto Aplicado Multiplataforma'):
    p(line, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

p('', space_after=0)
p('DOCUMENTAÇÃO DE ENTREGA', bold=True, size=14,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=0)
p('Sprints 1 a 4', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
p('', space_after=0)
p('Fortaleza — 2025', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18)

# ======================== 1. NOME DO PROJETO =======================
heading('1. Nome do projeto')
p('FisioTech — aplicativo mobile de gestão para clínicas e consultórios de fisioterapia, '
  'reunindo cadastro de pacientes e profissionais, agendamento e acompanhamento clínico de '
  'consultas e comunicação entre paciente e profissional em uma única base de código, '
  'distribuída como aplicação web e como aplicativo Android nativo.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)
p('Supervisora: Profª. Lyndainês Santos.')

# ========================== 2. INTEGRANTES =========================
heading('2. Integrantes')
build_table(['Nome do integrante', 'Matrícula'],
            [['Gabriel Raniere Barbosa Neri', '2412891'],
             ['Pedro Italo do Nascimento Costa', '2426040'],
             ['Marco Antonio Pessoa Noronha', '2427661'],
             ['Larissa Maria Félix de Oliveira', '2422783'],
             ['Caio Gabriel Rodrigues de Melo', '2422764']],
            [10.0, 5.0])

# ======================= 3. LINKS DO PROJETO =======================
heading('3. Links do projeto')
p('Repositórios de código e quadro de acompanhamento da equipe.')
links = DOC.add_table(rows=1, cols=2)
links.style = 'Table Grid'
links.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(['Recurso', 'Endereço']):
    cell_text(links.rows[0].cells[i], h, bold=True)
    shade(links.rows[0].cells[i], 'DEEAF6')
for nome, url in [('Repositório (back-end)', 'https://github.com/gabrielneriqa/fisiotech-back'),
                  ('Repositório (front-end)', 'https://github.com/gabrielneriqa/fisiotech-front'),
                  ('Project (board)', 'https://github.com/gabrielneriqa/fisiotech-back/projects')]:
    cells = links.add_row().cells
    cell_text(cells[0], nome)
    cells[1].text = ''
    par = cells[1].paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.space_before = Pt(2)
    hyperlink(par, url, 'Link')
fix_widths(links, [6.5, 8.5])

# ==================== 4. VISÃO GERAL DAS SPRINTS ===================
heading('4. Visão geral das sprints')
p('O desenvolvimento do FisioTech foi organizado em quatro sprints. A Sprint 1 estabeleceu a '
  'fundação técnica — estrutura do aplicativo, empacotamento Android e modelagem do banco. As '
  'Sprints 2, 3 e 4 construíram a API sobre essa base, avançando do controle de acesso para o '
  'autoatendimento do paciente e, por fim, para o ciclo completo do atendimento e a garantia '
  'de qualidade. Todos os itens listados neste documento encontram-se concluídos e '
  'incorporados à aplicação.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)

build_table(['Sprint', 'Foco', 'Itens', 'Status'],
            [['Sprint 1', 'Fundação do aplicativo e do banco de dados', '#1 a #4', 'Concluída'],
             ['Sprint 2', 'Autenticação, papéis de acesso e domínio clínico', '#5 a #11', 'Concluída'],
             ['Sprint 3', 'Autoatendimento do paciente e agendamento', '#12 a #18', 'Concluída'],
             ['Sprint 4', 'Ciclo do atendimento, qualidade e ajustes', '#19 a #26', 'Concluída']],
            [2.3, 8.2, 2.3, 2.2])

# ============================= SPRINT 1 ============================
heading('5. Sprint 1 — Fundação do aplicativo e do banco de dados', page_break=True)
p('A Sprint 1 concentrou a fundação técnica do projeto: a estrutura da aplicação Angular com '
  'navegação por papel de usuário, o empacotamento Android via Capacitor, a modelagem e criação '
  'do banco de dados com as entidades base e a produção da documentação de entrega da sprint.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)
build_table(['Item', 'Atividade', 'Status'],
            [['Fisio-Teach #1', 'Configurar Angular mobile-first com navegação por papel', 'Concluído'],
             ['Fisio-Teach #2', 'Integrar Capacitor e gerar o primeiro .apk', 'Concluído'],
             ['Fisio-Teach #3', 'Modelar e criar o banco com as entidades base', 'Concluído'],
             ['Fisio-Teach #4', 'Criação de documentação e entrega de sprint', 'Concluído']],
            [3.4, 9.2, 2.4])
figura(1, 1)

# ============================= SPRINT 2 ============================
heading('6. Sprint 2 — Autenticação, papéis de acesso e domínio clínico', page_break=True)
p('A Sprint 2 transformou a base modelada em uma API funcional e protegida. O ponto de partida '
  'foi a correção de duas falhas de segurança herdadas do esqueleto inicial — senha gravada em '
  'texto puro na edição e credenciais fixas no código. Em seguida foram implementados os três '
  'papéis de acesso (administrador, profissional e paciente), a entidade Profissional como dona '
  'dos próprios pacientes e o núcleo clínico do sistema: consulta, mensagem e avaliação.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)
build_table(['Item', 'Atividade', 'Status'],
            [['Fisio-Teach #5', 'Corrigir a senha gravada em texto puro e remover credenciais fixas no código', 'Concluído'],
             ['Fisio-Teach #6', 'Criar a entidade Profissional e vincular o paciente ao profissional responsável', 'Concluído'],
             ['Fisio-Teach #7', 'Implementar a autenticação e a autorização por papel de usuário', 'Concluído'],
             ['Fisio-Teach #8', 'Expor o endpoint de identificação do usuário autenticado', 'Concluído'],
             ['Fisio-Teach #9', 'Modelar Consulta, Mensagem e Avaliação com o registro clínico do atendimento', 'Concluído'],
             ['Fisio-Teach #10', 'Habilitar o CORS para o aplicativo empacotado com Capacitor', 'Concluído'],
             ['Fisio-Teach #11', 'Documentar o setup da API e criar os scripts de execução em modo de desenvolvimento', 'Concluído']],
            [3.4, 9.2, 2.4])
figura(2, 2)

# ============================= SPRINT 3 ============================
heading('7. Sprint 3 — Autoatendimento do paciente e agendamento', page_break=True)
p('Na Sprint 3 o paciente deixou de ser um registro gerenciado pelo profissional e passou a ser '
  'um usuário autônomo do aplicativo: criou a própria conta, encontrou profissionais por '
  'especialidade, consultou os horários livres na agenda e marcou a consulta sozinho. Do lado '
  'do profissional, a caixa de entrada unificada reuniu em uma única tela as conversas de todos '
  'os pacientes, e o administrador ganhou a gestão centralizada dos cadastros.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)
build_table(['Item', 'Atividade', 'Status'],
            [['Fisio-Teach #12', 'Criar a área de autoatendimento do paciente autenticado', 'Concluído'],
             ['Fisio-Teach #13', 'Entregar a caixa de entrada unificada de mensagens do profissional', 'Concluído'],
             ['Fisio-Teach #14', 'Adicionar a gestão de pacientes pelo administrador', 'Concluído'],
             ['Fisio-Teach #15', 'Implementar o autocadastro público de paciente', 'Concluído'],
             ['Fisio-Teach #16', 'Disponibilizar a busca de profissionais e a consulta de horários livres', 'Concluído'],
             ['Fisio-Teach #17', 'Permitir o agendamento da consulta pelo próprio paciente', 'Concluído'],
             ['Fisio-Teach #18', 'Permitir a edição de perfil e a troca de senha pelo paciente', 'Concluído']],
            [3.4, 9.2, 2.4])
figura(3, 3)

# ============================= SPRINT 4 ============================
heading('8. Sprint 4 — Ciclo do atendimento, qualidade e ajustes', page_break=True)
p('A Sprint 4 fechou o ciclo de vida do atendimento — cancelamento, remarcação e conversa '
  'vinculada a cada consulta — e concentrou o esforço de qualidade: correção da regra que '
  'permitia avaliar uma consulta ainda não realizada, cobertura dos services principais com '
  'testes de unidade e os ajustes de usabilidade levantados na validação com a equipe.',
  align=WD_ALIGN_PARAGRAPH.JUSTIFY)
build_table(['Item', 'Atividade', 'Status'],
            [['Fisio-Teach #19', 'Implementar o cancelamento e a remarcação de consulta', 'Concluído'],
             ['Fisio-Teach #20', 'Vincular as mensagens à consulta e exibir a conversa por atendimento', 'Concluído'],
             ['Fisio-Teach #21', 'Substituir o popup nativo de autenticação do navegador por resposta da própria API', 'Concluído'],
             ['Fisio-Teach #22', 'Corrigir a avaliação permitida antes de a consulta ser realizada', 'Concluído'],
             ['Fisio-Teach #23', 'Cobrir os services principais com testes de unidade', 'Concluído'],
             ['Fisio-Teach #24', 'Tornar a senha opcional na edição de profissional e de paciente', 'Concluído'],
             ['Fisio-Teach #25', 'Adicionar o autoatendimento de troca de senha para profissional e administrador', 'Concluído'],
             ['Fisio-Teach #26', 'Atualizar a documentação e realizar a entrega das sprints', 'Concluído']],
            [3.4, 9.2, 2.4])
figura(4, 4)

out = '/home/user/fisiotech-back/docs/FisioTech_Documentacao_Entrega_Sprints.docx'
DOC.save(out)
print('salvo:', out)
